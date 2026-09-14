#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import docx, sys
p = "/data/cy/shujuku/output/路面性能数智化平台总体设计与模块搭建方案（7域主体-12域全景·v1.1含标准数字化）.docx"
d = docx.Document(p)
print("paragraphs:", len(d.paragraphs), "tables:", len(d.tables))
print("---- 标题层级 ----")
for i, para in enumerate(d.paragraphs):
    t = para.text.strip()
    if not t:
        continue
    st = para.style.name
    if st.startswith("Heading") or st.startswith("标题") or t.startswith(("附录", "九、", "八、", "七、")):
        print(f"[{st}] {t[:80]}")
print("---- 表格首行（前 3 张） ----")
for k, tb in enumerate(d.tables[:3]):
    print(k, [c.text[:14] for c in tb.rows[0].cells])
print("---- 表格总数与末 3 张首行 ----")
for k, tb in enumerate(d.tables[-3:]):
    print(len(d.tables) - 3 + k, [c.text[:14] for c in tb.rows[0].cells])
txt = "\n".join(x.text for x in d.paragraphs)
for key in ["标准数字化", "能力等级", "本体", "SHACL", "保障架构", "协同", "附录一", "附录二", "44721"]:
    print(f"命中 {key}: {txt.count(key)}")
