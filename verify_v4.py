#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""校验 v1.2 报告结构：九章（9.1-9.7）+ 十章（10.1-10.6）+ 附录一/二 顺序与完整性"""
import re, sys
from docx import Document

DOCX = "/data/cy/shujuku/output/路面性能数智化平台总体设计与模块搭建方案（7域主体-12域全景·v1.2含标准数字化与对标分析）.docx"
d = Document(DOCX)
heads = [p.text.strip() for p in d.paragraphs if p.text.strip() and (p.style.name or "").startswith("Heading")]
doc_heads = [t for t in heads if re.match(r"^[一二三四五六七八九十]+、", t)]
sub = [t for t in heads if re.match(r"^(9|10)\.[0-9]+", t)]
apx = [t for t in heads if t.startswith("附录")]

print("=== 一级章节 ===")
for t in doc_heads:
    print(" ", t)
print("\n=== 9.x / 10.x 小节 ===")
for t in sub:
    print(" ", t)
print("\n=== 附录 ===")
for t in apx:
    print(" ", t)

ok = True
if not any("十、" in t for t in doc_heads):
    ok = False; print("!! 缺第十章")
for i in range(1, 8):
    if not any(t.startswith(f"9.{i}") for t in sub):
        ok = False; print(f"!! 缺 9.{i}")
for i in range(1, 7):
    if not any(t.startswith(f"10.{i}") for t in sub):
        ok = False; print(f"!! 缺 10.{i}")
if len(apx) < 2:
    ok = False; print("!! 附录不足两个")
# 顺序检查
order_key = []
for t in heads:
    if re.match(r"^九、", t): order_key.append(("九", heads.index(t)))
    if re.match(r"^十、", t): order_key.append(("十", heads.index(t)))
    if t.startswith("附录一"): order_key.append(("附录一", heads.index(t)))
    if t.startswith("附录二"): order_key.append(("附录二", heads.index(t)))
seq = [k for k, _ in sorted(order_key, key=lambda x: x[1])]
print("\n顺序:", " → ".join(seq))
if seq != ["九", "十", "附录一", "附录二"]:
    ok = False; print("!! 章节顺序异常")
print("\n表格数:", len(d.tables), "| 段落数:", len(d.paragraphs))
print("RESULT:", "OK ✅" if ok else "FAIL ❌")
sys.exit(0 if ok else 1)
