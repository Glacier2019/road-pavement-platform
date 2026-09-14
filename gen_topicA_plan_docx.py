# -*- coding: utf-8 -*-
"""生成《纯视觉道路线形曲率/坡度在线估计研究——论文大纲·技术路线·学生实施计划》docx
用法: PYTHONPATH="" /data/cy/shujuku/.docxvenv/bin/python gen_topicA_plan_docx.py
"""
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.top_margin = sec.bottom_margin = Mm(22)
sec.left_margin = sec.right_margin = Mm(25)

style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

def _set_font(run, size=12, bold=False, color=None, name='宋体'):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    if color:
        run.font.color.rgb = RGBColor(*color)

def h1(text):
    p = doc.add_paragraph()
    _set_font(p.add_run(text), 15, True, (0x1F, 0x3A, 0x5F), '黑体')
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(16), Pt(8)

def h2(text):
    p = doc.add_paragraph()
    _set_font(p.add_run(text), 13, True, (0x2C, 0x4A, 0x75), '黑体')
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(10), Pt(5)

def h3(text):
    p = doc.add_paragraph()
    _set_font(p.add_run(text), 12, True, (0x44, 0x44, 0x44), '黑体')
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(6), Pt(3)

def p(text, bold=False, size=12, indent=True):
    para = doc.add_paragraph()
    if indent:
        para.paragraph_format.first_line_indent = Pt(24)
    _set_font(para.add_run(text), size, bold)
    para.paragraph_format.space_after = Pt(4)
    para.paragraph_format.line_spacing = 1.3

def bullet(text, bold_prefix=None):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Pt(18)
    para.paragraph_format.space_after = Pt(3)
    para.paragraph_format.line_spacing = 1.3
    if bold_prefix:
        _set_font(para.add_run('· ' + bold_prefix), 12, True)
        _set_font(para.add_run(text), 12)
    else:
        _set_font(para.add_run('· ' + text), 12)

def table(headers, rows, widths=None, fontsize=10.5):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = 'Table Grid'
    for j, htxt in enumerate(headers):
        cell = t.rows[0].cells[j]
        cell.text = ''
        _set_font(cell.paragraphs[0].add_run(htxt), fontsize, True, name='黑体')
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.rows[i + 1].cells[j]
            cell.text = ''
            _set_font(cell.paragraphs[0].add_run(str(val)), fontsize)
    if widths:
        for j, w in enumerate(widths):
            for row in t.rows:
                row.cells[j].width = Mm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

# ============ 封面 ============
tp = doc.add_paragraph(); tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
_set_font(tp.add_run('纯视觉道路线形曲率/坡度在线估计研究'), 20, True, (0x1F, 0x3A, 0x5F), '黑体')
sp = doc.add_paragraph(); sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
_set_font(sp.add_run('—— 论文大纲 · 技术路线 · 学生实施计划（课题甲）——'), 14, True, (0x2C, 0x4A, 0x75), '黑体')
sp2 = doc.add_paragraph(); sp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
_set_font(sp2.add_run('以公路三维曲面方程为基础 · 路面起伏复用 RoadBEV · 在线估计借鉴 LingBot-Map 流式思路 · 几何作为约束'), 11, False, (0x66, 0x66, 0x66))
mp = doc.add_paragraph(); mp.alignment = WD_ALIGN_PARAGRAPH.CENTER
_set_font(mp.add_run('2026 年 9 月制定 · 面向 0 基础硕士研究生 · 周期 2026.09–2027.09（与 2025Y095 交付节点对齐）'), 10.5, False, (0x99, 0x99, 0x99))

# ============ 一、研究定位 ============
h1('一、研究定位（纯视觉方向确认）')
p('本课题确定为纯视觉技术路线，定义与边界如下：', bold=True)
bullet('传感器配置：车载单目或双目相机（视频输入）+ IMU/轮速（仅作尺度锚定与位姿辅助，不提供几何真值）。', bold_prefix='纯视觉定义——')
bullet('不依赖 LiDAR、高精地图、RTK/惯导精定位、先验点云。适用去图化（无图/轻图）与众包场景：低硬件成本、可大规模部署、可动态更新。', bold_prefix='排除项——')
bullet('路面起伏（高程场）由 RoadBEV 式纯视觉 BEV 重建获得（单目 1.83cm / 双目 0.50cm 高程误差水平，RSRD 公开数据集）；曲率/坡度估计以公路三维曲面方程（Darboux 标架几何不变量）为理论基础，借鉴 LingBot-Map 流式重建架构（桩号锚定+轨迹记忆+漂移校正）实现在线估计，JTG D20 规范与 GE 域设计先验作为几何约束。', bold_prefix='技术支柱——')
bullet('本课题的差异化卖点不是“在线重建”（RoadBEV 已做），而是“重建→线形参数化估计”的完整链路：输出曲率 κ(s)/坡度 G(s)/超高 E(s)/挠率 τ(s) 连续序列，并以滚动-振动模型耦合验证。', bold_prefix='定位声明——')

# ============ 二、论文大纲 ============
h1('二、论文大纲')
tp2 = doc.add_paragraph(); tp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
_set_font(tp2.add_run('建议题目：《基于公路三维曲面方程的纯视觉道路线形曲率与坡度在线估计方法研究》'), 12.5, True, (0x1F, 0x3A, 0x5F), '黑体')
p('学科：交通运输工程（道路与铁道工程方向）｜类型：工学硕士论文', size=11)

h2('第一章 绪论')
bullet('研究背景与意义：自动驾驶去图化/众包趋势；道路几何（曲率/坡度/超高/起伏）对行车安全与舒适的影响机理；悬架预瞄与线形复核的工程需求；公路基础设施数据库底座 C 端板块定位。')
bullet('国内外研究现状（四线并述+差距表）：①视觉坡度/曲率估计（IET-ITS 2019 等）；②单目 3D 车道检测与地图先验融合（Depth3DLane、PriorLane、P-MapNet）；③道路线形提取与重建（GPS/LiDAR/摄影测量：CACIE 2020 竖曲线参数估计、Springer 2023 综述、Camacho-Torregrosa 2015 噪声敏感性）；④路面高程重建（RoadBEV T-ITS 2024）与 3DGS 重建（RoGS/CurveGaussian）。')
bullet('研究内容与技术路线（对应第 2~5 章）；论文组织结构。')
bullet('开题差异化声明：对 RoadBEV（高程场 vs 线形参数输出）、对 3DGS 系（通用重建 vs 设计曲面解析参数化+规范约束）、对线形提取领域（离线测量源 vs 纯视觉在线众包源）。')

h2('第二章 公路三维曲面方程与线形参数化模型（理论基础）')
bullet('2.1 公路三维曲面描述：Darboux 标架下曲面几何不变量——测地曲率 κg、法曲率 κn、测地挠率 τg；平纵横耦合的曲面表达（圆曲线×纵坡→空间螺旋线，曲率受平纵共同影响）。')
bullet('2.2 线形要素参数化：直线/圆曲线/回旋线/竖曲线/超高渐变段的参数化（R、A、i1、i2、E(s)），与曲面方程参数的映射关系。')
bullet('2.3 估计问题定义：由高程场 E(x,y)（或点云/网格）反演曲面参数的不适定性分析；可辨识性与多解性讨论（开题前数值预研）。')
bullet('2.4 几何约束建模：JTG D20《公路路线设计规范》约束（最大纵坡、最小半径、超高渐变率、G² 线形连续）+ GE 域设计先验容差带建模（含设计-施工偏差处理）。')

h2('第三章 基于 RoadBEV 的路面高程场重建（纯视觉数据基础）')
bullet('3.1 RoadBEV 原理复现：BEV 体素、单目/双目两模型、RSRD 数据集（2800 对立体图像、3cm 网格、1.9m×5m 轮胎轨迹 ROI、点云 GT）。')
bullet('3.2 模型训练/微调与评测：高程误差、频带分析（区分线形尺度与微观起伏尺度）。')
bullet('3.3 尺度锚定：单目绝对尺度模糊的解决——已知车高基准面、轮速/IMU 融合、或双目起步。')
bullet('3.4 多帧拼接与 5m→50m 前瞻扩展：为第 4 章流式估计提供输入规范。')

h2('第四章 曲面参数化拟合与流式线形估计（核心创新章）')
bullet('4.1 流式状态估计框架（创新③）：桩号锚定与坐标系定义；帧间状态传播；轨迹记忆（已行驶路段曲面参数缓存）；重复经过路段的一致性校正（漂移抑制）——架构思路借鉴 LingBot-Map 的 Geometric Context Transformer 与轨迹记忆，但目标与表示不同。')
bullet('4.2 带约束的曲面参数化拟合（创新①②）：优化目标 = 高程数据项 + 曲面平滑项 + G² 连续性项 + JTG D20 规范硬约束项 + GE 域设计先验容差项；求解方法（scipy 序列二次规划 / 可微优化）；数值稳定性与初始化策略。')
bullet('4.3 线形参数求解：由拟合曲面参数解析/数值求解 κ(s)、G(s)、E(s)、τ(s) 连续序列，含置信度输出。')
bullet('4.4 消融实验：约束项逐项移除的精度-稳定性-合规性对比，证明每个约束的贡献（回应“为约束而约束”质疑）。')

h2('第五章 实验验证与应用分析')
bullet('5.1 RSRD 定量验证：线形参数估计误差 vs 点云 GT；与 RoadBEV+后处理（数值微分/B 样条拟合）、MLS 基线对比；消融结果汇总。')
bullet('5.2 G228 试验段实拍验证：估计线形 vs GE 域换算设计线形；竣工-设计偏差反演分析（工程应用叙事）。')
bullet('5.3 滚动-振动模型耦合验证（创新④）：估计误差经 1/4 车滚动-振动模型的受力/振动误差传播分析；曲面方程应用场景扩展。')
bullet('5.4 应用讨论：悬架预瞄（50~100m 前瞻）、线形一致性复核、底座 K 域众包回填；局限性（雨雾、遮挡、极端场景）。')

h2('第六章 结论与展望')
bullet('主要结论、创新点总结；展望：3DGS 全套方案（产业/专利版）、多车众包协同、与底座双 Copilot 联动。')

h2('参考文献与附录')
bullet('参考文献：RoadBEV（T-ITS 2024）、RoGS、CurveGaussian、AutoSplat、GeoSplat、LingBot-Map、Depth3DLane、PriorLane、P-MapNet、IET-ITS 2019 坡度估计、CACIE 2020 竖曲线重建、Springer 2023 线形提取综述、JTG D20 规范、导师论文（三维曲面滚动-振动）等（全部可核实）。')
bullet('附录：曲面参数化推导细节、优化器配置、数据集处理脚本说明。')

# ============ 三、技术路线图 ============
h1('三、技术路线图')
p('（深色底图，图中指标为目标值；M2/M4 为理论核心，M3 为架构创新，V3 为特色验证）', size=10.5)
try:
    doc.add_picture('/data/cy/shujuku/output/技术路线图-课题甲-纯视觉线形估计.png', width=Mm(158))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
except Exception as e:
    p('图片插入失败: %s' % e, size=10)

# ============ 四、学生实施计划 ============
h1('四、学生实施计划（0 基础，2026.09–2027.09）')

h2('4.1 总体时间线（与平台项目对齐）')
table(['阶段', '时间', '主线任务', '里程碑/产出'], [
    ['基础期', '2026.9–10', 'Python/PyTorch 集训 + RoadBEV 复现 + 微分几何入门', '跑通 RoadBEV 官方 demo；曲率噪声数值预研报告'],
    ['理论期 T1', '2026.11–12', '第 2 章公式推导 + 几何约束建模 + GE 域换算工具', '曲面参数化仿真验证；预研判据结论'],
    ['方法期 T2', '2027.1–4', '第 4 章流式估计 + 约束拟合 + 消融', '核心方法跑通；消融实验完整'],
    ['验证期 T3', '2027.5–7', 'RSRD 定量 + G228 实拍 + 振动模型耦合', '2027.5 中期检查；对比数据完整'],
    ['论文期', '2027.7–9', '论文写作 + 期刊投稿 + 预答辩', '2027.9 论文定稿送审；期刊初稿投出'],
], widths=[20, 20, 62, 60])

h2('4.2 分阶段任务详解（现在做什么 / 学什么 / 下一步）')

h3('阶段 0 · 基础期（2026.9–10，6 周）')
p('现在做什么：', bold=True)
bullet('W1-2：通读导师论文《沿公路三维曲面车轮滚动-振动受力分布特征研究》（理解 Darboux 标架、几何不变量、双激励模型）；精读 RoadBEV 论文与 RSRD 数据集说明；浏览 zidongjiashi 文件夹 4 份方案文档（重点：参考论文清单）；建术语表（曲率/坡度/超高/挠率/回旋线/桩号/IRI）。')
bullet('W3-4：PyTorch 集训（张量、Dataset/DataLoader、训练循环、CNN 分类/回归、GPU 训练注意事项）；Python 数值计算（numpy/scipy/matplotlib）；Linux/Git 基础。')
bullet('W5-6：复现 RoadBEV 推理流程——在 RSRD 测试集上跑通官方 checkpoint，可视化 BEV 高程图与误差图；同时学微分几何入门（曲线曲率/挠率、Frenet 标架、曲面第一/第二基本形式、Darboux 标架概念，用课本+讲义，不需全懂，够用即可）。')
p('学什么：', bold=True)
bullet('Python/PyTorch 深度学习基础；微分几何入门（聚焦曲线与曲面的曲率类概念）；相机模型与坐标变换（针孔模型、外参内参、BEV 投影）；Git 协作规范。')
p('验收标准：', bold=True)
bullet('RoadBEV demo 跑通并输出误差可视化图；完成“高程场→数值差分曲率→噪声量级”最小实验（用 RSRD 一张图即可），写 1 页预研结论。')

h3('阶段 1 · 理论期 T1（2026.11–12）★ 开题前关键判据')
p('现在做什么：', bold=True)
bullet('第 2 章核心推导：将公路线形要素（直线/圆曲线/回旋线/竖曲线/超高）统一到三维曲面方程参数化框架，推导几何不变量→线形参数的映射公式（Mathematica/scipy 符号+数值验证）。')
bullet('★ 曲率噪声数值预研（论文立论判据）：在 RSRD 真实高程场+GT 上，对比三条路线——① 直接数值微分求曲率；② B 样条拟合后求曲率；③ 带规范约束的曲面参数化拟合求曲率（本课题路线）——给出三者的曲率误差量级对比与结论：约束/参数化能否把曲率误差压到可用水平（目标：与设计值/点云 GT 偏差可比，如 κ 误差 < 2×10⁻⁴ m⁻¹ 或给出可达量级）。预研结论决定论文定位与发表档位。')
bullet('GE 域换算工具 v1：把 G228 设计文件（平曲线要素表/纵断面设计/超高渐变表）换算为 κ(s)/G(s)/E(s) 函数库（与课题乙共用，可向平台底座组同学提接口需求）。')
bullet('3.3 尺度锚定方案设计：已知车高基准面 / 轮速积分 / 双目选型的定量对比。')
p('学什么：', bold=True)
bullet('微分几何深化（曲面方程、第一/第二基本形式、测地曲率法曲率测地挠率计算）；数值优化入门（最小二乘、带约束优化、SQP/ADMM 概念）；scipy.optimize 与 PyTorch 自动微分实战。')
p('验收标准：', bold=True)
bullet('预研报告（曲率误差对比三路线结论）；第 2 章公式推导文档 v1；GE 域换算工具 v1 入库（Git）。')

h3('阶段 2 · 方法期 T2（2027.1–4）核心实现')
p('现在做什么：', bold=True)
bullet('1 月：第 3 章收尾——RoadBEV 微调（RSRD 全量训练/微调，记录高程误差与频带特征）；多帧拼接模块 v1（帧间位姿由 IMU/轮速+视觉里程计给出）。')
bullet('2 月：4.2 带约束曲面拟合优化器 v1（先离线 batch 版：数据项+平滑+G²+规范硬约束+设计先验），在 RSRD 上跑通并记录收敛行为。')
bullet('3 月：4.1 流式状态估计 v1——桩号锚定、帧间传播、轨迹记忆、重复路段一致性校正（借鉴 LingBot-Map 架构思路实现轻量版：状态=局部曲面参数，记忆=已过桩号段参数缓存，校正=重访段的加权平均+残差检测）。')
bullet('4 月：4.3 线形参数输出模块（κ/G/E/τ 序列+置信度）；4.4 消融实验（约束逐项移除，产出精度/平滑度/合规性对比表）。')
p('学什么：', bold=True)
bullet('优化理论与可微编程；卡尔曼/因子图类状态估计基础（可选，够用即可）；实验方法论（消融设计、指标选择、误差传播分析）。')
p('验收标准：', bold=True)
bullet('2027.3 全链路可演示（视频输入→线形参数流输出）；消融表完整；每周 Git 提交即开发日志。')

h3('阶段 3 · 验证期 T3（2027.5–7）')
p('现在做什么：', bold=True)
bullet('5 月：RSRD 定量实验收尾——线形参数误差 vs 点云 GT；与 RoadBEV+后处理、B 样条、MLS 基线完整对比；中期检查材料（2027.5）。')
bullet('6 月：G228 实拍采集与验证（与平台组协调试验段通行与相机安装；估计线形 vs GE 域设计线形；竣工-设计偏差分析）。')
bullet('7 月：滚动-振动模型耦合——将估计的 κ/G/E 序列输入 1/4 车模型，对比“真值输入 vs 估计输入”的受力/振动差异，做误差传播分析（导师协助建模细节）；论文初稿 60%（第 1~3 章+方法章）。')
p('验收标准：', bold=True)
bullet('对比实验全部完成；误差传播分析结论明确；初稿 60%。')

h3('阶段 4 · 论文期（2027.7–9）')
p('现在做什么：', bold=True)
bullet('论文全稿写作与反复修改；期刊论文提炼与投稿（见 4.5）；预答辩准备（2027.9）；结题材料（实现说明/测试报告）。')

h2('4.3 考核与节奏')
table(['机制', '要求'], [
    ['契约先行', '第 2 章公式、数据集接口、GE 域函数库格式为“唯一事实源”，变更走工单由导师把关'],
    ['两周里程碑', '每两周一个可演示/可量化进展；每周 30 分钟例会（与平台组例会合并）'],
    ['Git 日志', 'commit 即开发日志；预研报告/推导文档/实验脚本全部入库'],
    ['写作训练', '每阶段产出 1 份技术文档（预研报告、公式推导、消融报告），即论文素材'],
], widths=[30, 132])

h2('4.4 风险与对策')
table(['风险', '等级', '对策'], [
    ['曲率噪声压不住（预研判据不通过）', '高', '开题前即暴露：降级为“高程重建+竣工线形反演应用”叙事（二档期刊），或转向悬架预瞄应用；不硬扛'],
    ['深度学习入门慢（第 3 章 RoadBEV 训练）', '中', '基础期 PyTorch 集训前置；训练范围限微调；与平台组共享 GPU 资源错峰'],
    ['单目尺度模糊', '中', '尺度锚定方案开题前定量对比（车高基准面/轮速/双目）；必要时双目起步'],
    ['前瞻距离不足（RoadBEV ROI 5m）', '中', '第 4 章流式拼接是正解；若超时，降级为“多帧拼接+滤波”并保留消融'],
    ['与平台组接口耦合（GE 域工具）', '低', '先与课题乙共用同一换算工具；接口契约先行'],
    ['发表进度风险', '中', '预研通过后 2027.3 投会议短文/期刊初稿；主论文 2027.9 前投出（见 4.5）'],
], widths=[70, 14, 78])

h2('4.5 发表计划')
table(['档位', '目标', '投稿时间', '条件'], [
    ['冲', 'Computer-Aided Civil and Infrastructure Engineering（IF≈11）/ TRC / IEEE T-ITS（IF≈8）', '2027.6–9 投主论文', '预研判据通过 + 消融完整 + RSRD 定量赢基线'],
    ['稳', 'IEEE T-IV / Automation in Construction / ESWA', '2027.6–9 备选', '应用叙事强化（竣工反演+悬架预瞄）'],
    ['保底', '中国公路学报 / 交通运输工程学报；或会议（IV/ITSC/TRB）', '2027.3 短文', '预研判据不理想时启用'],
], widths=[18, 68, 30, 46])
p('预研结论决定投稿档位：若“带约束参数化”曲率误差达可用水平（相对设计值/GT 偏差可比）→ 冲一档；否则转应用叙事走二档。', size=11, bold=True)

# ============ 五、与底座/课题乙的协同 ============
h1('五、与底座及课题乙的协同')
bullet('GE 域：本课题是 GE 域“设计文件→κ/G/E 函数库”的第一个真实消费方，换算工具与课题乙共用，由底座数据组同学牵头。')
bullet('K 域回流：本课题估计结果（众包线形修正）设计为底座 K 域的数据来源之一，回填后持续更新 GE 域先验容差带。')
bullet('课题乙：共享 GE 工具与 G228 实拍数据（课题乙用几何约束做换道轨迹，本课题用几何约束做估计——两课题各消费 H 域一半）。')
bullet('悬架预瞄/数字孪生/3DGS 全套：论文展望章 1~2 页；zidongjiashi 文档的 3DGS 方案作为后续专利与产业版路线，不作为本论文主体。')

end = doc.add_paragraph(); end.alignment = WD_ALIGN_PARAGRAPH.CENTER
_set_font(end.add_run('—— 任务一经分配即冻结，照图施工，设计变更一律由导师把关 ——'), 11, False, (0x88, 0x88, 0x88))

out = '/data/cy/shujuku/output/纯视觉线形估计-论文大纲技术路线与学生实施计划（课题甲）.docx'
doc.save(out)
print('OK ->', out)
