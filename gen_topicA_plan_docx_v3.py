# -*- coding: utf-8 -*-
"""生成《课题甲 v3：3DGS+弧长模型+流式在线估计》docx
用法: PYTHONPATH="" /data/cy/shujuku/.docxvenv/bin/python gen_topicA_plan_docx_v3.py
"""
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.top_margin = sec.bottom_margin = Mm(22)
sec.left_margin = sec.right_margin = Mm(25)

style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

def _set_font(run, size=12, bold=False, color=None, name='宋体'):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    if color:
        run.font.color.rgb = RGBColor(*color)

def h1(text):
    p = doc.add_paragraph()
    _set_font(p.add_run(text), 15, True, (0x1F, 0x3A, 0x5F), '黑体')
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(16), Pt(8)

def h2(text):
    p = doc.add_paragraph()
    _set_font(p.add_run(text), 13, True, (0x2C, 0x4A, 0x75), '黑体')
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(10), Pt(5)

def h3(text):
    p = doc.add_paragraph()
    _set_font(p.add_run(text), 12, True, (0x44, 0x44, 0x44), '黑体')
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(6), Pt(3)

def p(text, bold=False, size=12, indent=True):
    para = doc.add_paragraph()
    if indent:
        para.paragraph_format.first_line_indent = Pt(24)
    _set_font(para.add_run(text), size, bold)
    para.paragraph_format.space_after = Pt(4)
    para.paragraph_format.line_spacing = 1.3

def bullet(text, bold_prefix=None):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Pt(18)
    para.paragraph_format.space_after = Pt(3)
    para.paragraph_format.line_spacing = 1.3
    if bold_prefix:
        _set_font(para.add_run('· ' + bold_prefix), 12, True)
        _set_font(para.add_run(text), 12)
    else:
        _set_font(para.add_run('· ' + text), 12)

def table(headers, rows, widths=None, fontsize=10.5):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = 'Table Grid'
    for j, htxt in enumerate(headers):
        cell = t.rows[0].cells[j]
        cell.text = ''
        _set_font(cell.paragraphs[0].add_run(htxt), fontsize, True, name='黑体')
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.rows[i + 1].cells[j]
            cell.text = ''
            _set_font(cell.paragraphs[0].add_run(str(val)), fontsize)
    if widths:
        for j, w in enumerate(widths):
            for row in t.rows:
                row.cells[j].width = Mm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

# ============ 封面 ============
tp = doc.add_paragraph(); tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
_set_font(tp.add_run('3DGS + 弧长线形模型 + 流式在线估计（课题甲 v3）'), 20, True, (0x1F, 0x3A, 0x5F), '黑体')
sp = doc.add_paragraph(); sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
_set_font(sp.add_run('—— SCI 论文题目 · 重构大纲 · 技术路线 · 可行性 · 硕士论文大纲 · 学生实施计划 ——'), 13, True, (0x2C, 0x4A, 0x75), '黑体')
sp2 = doc.add_paragraph(); sp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
_set_font(sp2.add_run('核心创新：3DGS 路面重建（载体） + 弧长参数化三维线形模型（表示层/先验） + LingBot-Map 式流式在线估计（几何约束） · 不依赖 RoadBEV · 输出设计语义参数并接车辆动力学预判'), 10.5, False, (0x66, 0x66, 0x66))
mp = doc.add_paragraph(); mp.alignment = WD_ALIGN_PARAGRAPH.CENTER
_set_font(mp.add_run('2026 年 9 月制定 · 面向 0 基础硕士研究生 · 周期 2026.09–2027.09'), 10.5, False, (0x99, 0x99, 0x99))

# ============ 一、核心创新与可行性 ============
h1('一、核心创新定稿与可行性分析（skills 深度分析）')

h2('1.1 核心创新一句话')
p('以 3D Gaussian Splatting 路面重建为视觉载体，把博士论文《公路几何设计三维线形评价模型研究》的「基于弧长积分的公路三维几何模型」（准线-直母线直纹面；直线/圆曲线/回旋线/竖曲线/超高分段；Darboux 标架几何不变量 κg/κn/τg）作为表示层与先验，借鉴阿里 LingBot-Map 的流式重建架构（桩号锚定/滑窗参考/轨迹记忆/分页缓存抗漂移）实现在线估计，把 3DGS 稠密几何精化为**设计语义参数**（R、A、i、E(s)、κ/G/τ 序列），并以车辆动力学模型验证其对运行安全舒适性预判的合理性。', bold=True)

h2('1.2 LingBot-Map 可借鉴点（精确到模块，只借思路不挪用模型）')
table(['LingBot-Map 模块（蚂蚁，Apache-2.0）', '借鉴到本课题', '落地场景'], [
    ['Geometric Context Transformer：坐标锚定+稠密几何线索+长程漂移校正统一框架', '桩号锚定 + 滑窗参考（已重建/记忆窗）+ 重复路段一致性校正', '3DGS 重建结果按桩号组织；滑窗内融合多视角几何线索'],
    ['轨迹记忆（trajectory memory）', '已行驶路段曲面/线形参数缓存，重访时加权融合+残差检测', '众包多趟采集融合（漂移抑制+一致性提升）'],
    ['paged KV cache 长序列（>10000 帧）', '按桩号分页缓存重建/估计状态，长行程内存可控', '车载长行程连续估计（对应底座 K 域众包）'],
    ['feed-forward 流式推理（~20FPS）', '帧级增量更新局部曲面状态，不 batch 离线', '在线估计的实时性需求'],
    ['VGGT 基础模型可加载双向推理（stage-1）', '（可选）深度/几何基础模型做稠密几何先验蒸馏', '优化 3DGS 初始化/稠密化'],
], widths=[56, 46, 60])
p('注意：LingBot-Map 本身是通用场景重建（室内为主），不输出曲率/坡度——借鉴的是「流式+抗漂移+记忆」架构思路，不是直接复用其模型；开题必须这样表述。', size=11)

h2('1.3 赛道判定：语义参数估计赛道 vs 重建赛道（不构成创新性撞车）')
table(['方法', '赛道', '输出', '局限（本方案补的）'], [
    ['CurveGaussian (ICCV 2025)', '重建（通用轮廓曲线）', '贝塞尔曲线坐标', '参数无工程语义；无规范约束；无动力学；输入为边缘图'],
    ['RoGS (arXiv 2405.14342)', '重建（道路路面网格）', '高程网格', '曲率后处理；无设计先验/超高/纵坡语义；无动力学验证；无流式在线'],
    ['GTLR-GS (arXiv 2603.23192)', '重建（室内 LiDAR 正则）', '渲染优化几何', '曲率细分渲染导向；非道路场景'],
    ['LaneCPP (CVPR 2024)', '感知（3D 车道线）', '连续 3D 曲线', '通用物理先验；不输出设计要素参数；无文件级先验'],
    ['LingBot-Map (蚂蚁 2026)', '流式场景重建', '3D 点云', '通用重建，不输出线形语义参数'],
    ['**本方案**', '**线形语义参数估计（流式）**', '**R/A/i/E(s)+κ/G/τ 序列**', '**设计语义输出+JTG D20 约束+车辆动力学预判+流式在线（他法所无）**'],
], widths=[40, 34, 30, 58])
p('结论：现有方法全部是"重建导向"；本方案是"语义导向"且加"流式在线"。部件级先例不影响组合创造性，开题必须写差异化声明：对 CurveGaussian/RoGS/GTLR-GS/LingBot-Map 逐条"重建导向 vs 语义导向"。', size=11)

h2('1.4 三实证清单（无实证=claim）')
bullet('E1 精化增益：「3DGS 重建 + 模型精化」vs「3DGS 直接重建/后处理求曲率」，R/A/i 设计参数精度提升（消融）。')
bullet('E2 可辨识性：含噪几何反演模型参数的唯一性与初值敏感性（开题前 2~4 周数值预研，三问）。')
bullet('E3 语义/物理合理性：JTG D20 合规率；设计-施工偏差检出；输入滚动-振动模型后与实测振动/受力吻合度。')

h2('1.5 风险评级与对策（v3 关键：3DGS 进主干的权衡）')
table(['风险', '等级', '对策'], [
    ['3DGS 工程量大（0 基础 12 个月）', '中高', '复用 RoGS 等开源方案（COLMAP+场景优化）；不重写渲染器；3DGS 限定为观测输入不强行可微精化'],
    ['3DGS 部件撞车（CurveGaussian/RoGS/GTLR-GS）', '中', '语义赛道差异化声明；不把"3DGS 本身"当卖点'],
    ['LingBot-Map 借鉴被质疑"拼装"', '中', '只借流式抗漂移架构思想，锁桩号锚定=GE 域闭环叙事；V2 专门验证长序列/多趟增益'],
    ['要素切分+参数估计初值/可辨识性', '中高', 'E2 数值预研开题前完成，结论决定档位'],
    ['单目尺度模糊 / 3DGS 尺度', '中', '3DGS 尺度由轨迹/轮速初始化；必要时双目/环视'],
], widths=[62, 14, 86])
p('总体评级：B+（可开题，3DGS 进主干后工程风险上升，但开源复用+语义叙事可对冲；三实证齐备可冲一档期刊）。', size=12, bold=True)

# ============ 二、SCI 论文 ============
h1('二、高水平 SCI 论文（题目 + 大纲 + 投稿定位）')

h2('2.1 推荐题目（主推 v3）')
tp2 = doc.add_paragraph(); tp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
_set_font(tp2.add_run('3D Gaussian Splatting-Based Road Curvature and Grade Estimation via Arc-Length-Parameterized Highway Alignment Model with Streaming Memory-aided Refinement'), 11.5, True, (0x1F, 0x3A, 0x5F), 'Times New Roman')
tp3 = doc.add_paragraph(); tp3.alignment = WD_ALIGN_PARAGRAPH.CENTER
_set_font(tp3.add_run('中文：基于 3D 高斯泼溅与弧长参数化公路三维线形模型的流式记忆增强道路曲率与坡度估计'), 12, True, (0x2C, 0x4A, 0x75), '黑体')

h2('2.2 论文大纲（SCI 六章）')
h3('Chapter 1 Introduction')
bullet('Background: 去图化/众包自动驾驶；道路几何对安全舒适影响；悬架预瞄与线形复核。')
bullet('Related work 五线+差距表：视觉坡度/曲率（IET-ITS 2019）；3D 车道检测与先验融合（LaneCPP/Depth3DLane/PriorLane）；道路线形提取重建（CACIE 2020）；3DGS 重建（RoGS/CurveGaussian/AutoSplat/GTLR-GS/DrivingForward）；流式重建（LingBot-Map）。')
bullet('Contribution（三创新）：①3DGS 重建载体+弧长参数化模型表示层；②流式记忆增强在线估计（借鉴 LingBot-Map）；③车辆动力学验证。')
h3('Chapter 2 Arc-Length-Parameterized Highway Alignment Model（创新①）')
bullet('2.1 直纹面曲面模型；2.2 平/纵/横分段参数化；2.3 Darboux 标架几何不变量；2.4 规范约束体系；2.5 估计问题定义与可辨识性。')
h3('Chapter 3 3DGS Scene Representation and Constrained Fusion（创新①/②）')
bullet('3.1 3DGS 路面重建（RoGS 复用，SfM 初始化）；3.2 3DGS 尺度锚定（轨迹/轮速）；3.3 要素切分+参数估计（R/A/i/E）；3.4 模型-数据融合精化（数据+平滑+G²+JTG D20+设计先验）。')
h3('Chapter 4 Streaming Memory-aided Online Estimation（创新②）')
bullet('4.1 桩号锚定的状态表示与滑窗参考；4.2 轨迹记忆与重访漂移校正；4.3 分页缓存长序列推理；4.4 帧级增量更新与实时性。')
h3('Chapter 5 Experiments（E1/E2/E3）')
bullet('5.1 合成+G228：R/A/i vs 设计真值；「3DGS+模型精化」vs「3DGS 直接重建/后处理」消融。')
bullet('5.2 长序列流式：单趟 vs 多趟重访，漂移量化+校正后一致性（V2）。')
bullet('5.3 初值鲁棒性（E2）；5.4 车辆动力学验证：Darboux→速度规划/滚动-振动模型，误差传播（E3）。')
bullet('5.5 工程应用：JTG D20 合规率、竣工偏差检出。')
h3('Chapter 6 Conclusion and Future Work')
bullet('结论；展望：端到端可微约束入 3DGS 渲染（产业版）、多车众包、底座 K 域回流。')

h2('2.3 投稿定位')
table(['档位', '目标', '条件'], [
    ['冲', 'IEEE T-ITS / TRC / CACIE（IF≈8-11）', '三实证齐备 + 定量赢基线 + 流式长序列增益明显'],
    ['稳', 'IEEE T-IV / Automation in Construction / ESWA', '应用叙事强化（竣工复核+合规率）'],
    ['保底', '中国公路学报 / 交通运输工程学报 / IV·ITSC 会议', '预研判据不理想时启用'],
], widths=[18, 80, 64])

# ============ 三、技术路线图 ============
h1('三、技术路线图（v3）')
p('（深色底图；M1-M4 为核心创新，V2 为流式/漂移校正专项验证；图中指标为目标值）', size=10.5)
try:
    doc.add_picture('/data/cy/shujuku/output/技术路线图-课题甲-3DGS流式线形估计-v3.png', width=Mm(158))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
except Exception as e:
    p('图片插入失败: %s' % e, size=10)

# ============ 四、硕士论文大纲 ============
h1('四、硕士学位论文大纲（正式版）')
p('题目：《基于 3D 高斯泼溅与弧长参数化公路三维线形模型的纯视觉道路曲率与坡度估计研究》', size=11)
h3('第一章 绪论')
bullet('1.1 研究背景与意义；1.2 国内外研究现状（五线+差距）；1.3 研究内容与技术路线；1.4 论文组织结构。')
h3('第二章 弧长参数化公路三维线形模型与语义表示')
bullet('2.1 公路三维曲面（直纹面）模型；2.2 平/纵/横分段参数化；2.3 Darboux 标架几何不变量；2.4 规范约束体系；2.5 估计问题定义与可辨识性。')
h3('第三章 基于 3DGS 的路面三维重建与尺度锚定')
bullet('3.1 3DGS 原理与 RoGS 复现；3.2 道路场景适配（SfM 初始化/路面高斯/动态物体滤除）；3.3 尺度锚定（轨迹/轮速/双目）；3.4 稠密几何输出与误差分析。')
h3('第四章 基于弧长线形模型的语义参数估计')
bullet('4.1 要素切分与参数估计框架；4.2 带约束的参数化拟合（数据+平滑+G²+JTG D20+设计先验）；4.3 初值策略与全局-局部优化；4.4 消融实验。')
h3('第五章 流式在线估计与记忆增强')
bullet('5.1 桩号锚定状态表示与滑窗参考；5.2 轨迹记忆与重访漂移校正；5.3 分页缓存长序列推理；5.4 帧级增量更新与实时性；5.5 长序列/多趟融合实验。')
h3('第六章 实验验证与应用分析')
bullet('6.1 合成+G228 定量验证（R/A/i vs 真值）；6.2 与直接重建/后处理基线对比；6.3 车辆动力学耦合验证（误差传播）；6.4 工程应用（合规率/竣工复核）。')
h3('第七章 结论与展望')
bullet('7.1 主要结论；7.2 创新点；7.3 展望。')
p('（含参考文献与附录：公式推导、优化器配置、数据处理脚本说明）', size=11)

# ============ 五、学生实施计划 ============
h1('五、学生实施计划（0 基础，2026.09–2027.09）')

h2('5.1 总体时间线')
table(['阶段', '时间', '主线任务', '里程碑/产出'], [
    ['基础期', '2026.9–10', 'Python/PyTorch 集训 + 博士论文模型 + 3DGS/RoGS 复现', 'RoGS demo 跑通；模型理解文档'],
    ['理论期 T1', '2026.11–12', '第 2 章推导 + ★E2 预研 + GE 域换算工具', '预研报告（决定档位）'],
    ['方法期 T2', '2027.1–4', '第 4 章语义参数估计 + 第 5 章流式估计 + 消融', '2027.3 全链路演示；消融表完整'],
    ['验证期 T3', '2027.5–7', '合成+G228 定量 + 长序列/多趟 + 动力学耦合', '2027.5 中期；对比数据完整；初稿 60%'],
    ['论文期', '2027.7–9', '硕论定稿 + SCI 提炼投稿 + 预答辩', '2027.9 送审；SCI 初稿投出'],
], widths=[20, 20, 62, 60])

h2('5.2 分阶段任务详解')
h3('阶段 0 · 基础期（2026.9–10，6 周）')
p('现在做什么：', bold=True)
bullet('W1-2：精读博士论文第 2 章（直纹面模型/Darboux/约束）；通读 RoGS 论文（3DGS 道路重建）、CurveGaussian（曲线高斯）、LingBot-Map 开源仓库（架构/README/HF 卡）。')
bullet('W3-4：PyTorch 集训；COLMAP/SfM 基础；3DGS 原理（高斯参数/可微渲染/稠密化）；Linux/Git。')
bullet('W5-6：复现 RoGS 道路重建 demo（合成或公开场景），可视化路面高斯/高程；用 Mathematica/scipy 复算博士论文 2.5 节几何公式。')
p('学什么：', bold=True)
bullet('PyTorch；3DGS 原理与渲染管线；COLMAP；微分几何入门；公路线形要素（JTG D20）。')
p('验收：RoGS demo 跑通；公式复算通过。', bold=True)

h3('阶段 1 · 理论期 T1（2026.11–12）★ 开题前判据')
p('现在做什么：', bold=True)
bullet('第 2 章核心推导：直纹面参数化→Darboux 不变量→R/A/i/E 映射公式（符号+数值验证）。')
bullet('★ E2 数值预研（2~4 周，论文立论判据）：在合成/仿真数据（已知 R/A/i 的曲面）上做「要素切分+参数估计」，三问：①含噪几何下要素分段与参数能否恢复；②多初值是否收敛同一解；③vs 样条/多项式基线增益。结论三档：可辨识+增益→冲一档；要素可切分但参数不稳→竣工反演叙事二档；仅符号/量级→悬架预瞄应用。')
bullet('GE 域换算工具 v1：G228 设计文件→κ(s)/G(s)/E(s) 函数库（与课题乙共用）。')
p('学什么：', bold=True)
bullet('数值优化（带约束/全局搜索）；混合整数优化（切分点）；scipy/PyTorch 自动微分。')
p('验收：预研报告（三问结论）；推导文档 v1；GE 工具 v1 入库。', bold=True)

h3('阶段 2 · 方法期 T2（2027.1–4）核心实现')
p('现在做什么：', bold=True)
bullet('1 月：RoGS 道路重建复现完整版（G228/合成场景）；3DGS 尺度锚定（轨迹/轮速）。')
bullet('2 月：要素切分+参数估计 v1（初值策略+全局-局部优化），3DGS 几何上跑通并记录收敛。')
bullet('3 月：带约束拟合完整版（数据+平滑+G²+JTG D20+设计先验）；全链路演示（2027.3）。')
bullet('4 月：流式在线估计 v1（桩号锚定+滑窗+轨迹记忆+漂移校正）；消融实验（约束逐项移除：精度/平滑度/合规性对比）。')
p('学什么：', bold=True)
bullet('优化理论与可微编程；状态估计基础（卡尔曼/因子图概念）；3DGS 稠密化/正则化；实验方法论。')
p('验收：全链路可演示（3DGS→线形参数流）；消融表完整；每周 Git 提交。', bold=True)

h3('阶段 3 · 验证期 T3（2027.5–7）')
p('现在做什么：', bold=True)
bullet('5 月：合成+G228 定量收尾（R/A/i vs 设计真值；vs 3DGS 直接重建/后处理基线）；中期检查材料。')
bullet('6 月：长序列流式验证（单趟 vs 多趟重访，漂移量化+校正后一致性）；G228 实拍采集（多趟）。')
bullet('7 月：车辆动力学耦合（Darboux→速度规划/滚动-振动模型，误差传播，与实测对比）；论文初稿 60%。')
p('验收：对比实验完整；E3 结论明确；初稿 60%。', bold=True)

h3('阶段 4 · 论文期（2027.7–9）')
p('硕论全稿与修改；SCI 提炼与投稿（见 2.3）；预答辩（2027.9）；结题材料（实现说明/测试报告）。')

h2('5.3 考核与节奏')
table(['机制', '要求'], [
    ['契约先行', '第 2 章公式、数据集接口、GE 域函数库为唯一事实源；变更走工单'],
    ['两周里程碑', '每两周可演示/可量化进展；每周 30 分钟例会'],
    ['Git 日志', 'commit 即开发日志'],
    ['写作训练', '每阶段 1 份技术文档（预研/推导/消融报告）=论文素材'],
], widths=[30, 132])

h2('5.4 风险与对策')
table(['风险', '等级', '对策'], [
    ['E2 预研判据不通过', '高', '开题前暴露：降级为竣工线形反演应用（二档）或悬架预瞄应用；不硬扛'],
    ['3DGS 工程量超期', '中高', '开源复用（RoGS）不做可微精化；必要时退回轻量观测（v2 路线）保毕业'],
    ['LingBot-Map 借鉴质疑', '中', '只借流式抗漂移思想；V2 专项验证长序列增益'],
    ['单目尺度模糊 / 3DGS 尺度', '中', '轨迹/轮速初始化；双目/环视备选'],
    ['发表进度', '中', '预研通过后 2027.3 投会议/短文；主论文 2027.9 前投出'],
], widths=[62, 14, 86])

h2('5.5 发表计划')
table(['档位', '目标', '投稿时间', '条件'], [
    ['冲', 'IEEE T-ITS / TRC / CACIE（IF≈8-11）', '2027.6–9 主论文', 'E2 预研通过 + 消融完整 + 定量赢基线 + 流式增益'],
    ['稳', 'IEEE T-IV / Automation in Construction / ESWA', '2027.6–9 备选', '应用叙事强化'],
    ['保底', '中国公路学报 / 交通运输工程学报；IV/ITSC 会议', '2027.3 短文', '预研判据不理想时启用'],
], widths=[18, 68, 30, 46])

# ============ 六、协同 ============
h1('六、与底座及课题乙的协同')
bullet('GE 域：本课题是"设计文件→κ/G/E 函数库"第一消费方，工具与课题乙共用。')
bullet('K 域回流：流式多趟估计结果（众包线形修正）设计为底座 K 域数据源，回填后更新 GE 域先验容差带。')
bullet('课题乙：共享 GE 工具与 G228 实拍（乙用几何约束做换道轨迹，甲用几何约束做估计——各消费 H 域一半）。')
bullet('数字孪生/悬架预瞄：论文展望章；与底座孪生板块衔接。')

end = doc.add_paragraph(); end.alignment = WD_ALIGN_PARAGRAPH.CENTER
_set_font(end.add_run('—— 任务一经分配即冻结，照图施工，设计变更一律由导师把关 ——'), 11, False, (0x88, 0x88, 0x88))

out = '/data/cy/shujuku/output/课题甲v3-3DGS流式线形估计-SCI论文大纲技术路线可行性硕论大纲与学生实施计划.docx'
doc.save(out)
print('OK ->', out)