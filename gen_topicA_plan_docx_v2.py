# -*- coding: utf-8 -*-
"""生成《课题甲 v2：SCI论文题目·重构大纲·技术路线·可行性·硕士论文大纲·学生实施计划》docx
用法: PYTHONPATH="" /data/cy/shujuku/.docxvenv/bin/python gen_topicA_plan_docx_v2.py
"""
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.top_margin = sec.bottom_margin = Mm(22)
sec.left_margin = sec.right_margin = Mm(25)

style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

def _set_font(run, size=12, bold=False, color=None, name='宋体'):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    if color:
        run.font.color.rgb = RGBColor(*color)

def h1(text):
    p = doc.add_paragraph()
    _set_font(p.add_run(text), 15, True, (0x1F, 0x3A, 0x5F), '黑体')
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(16), Pt(8)

def h2(text):
    p = doc.add_paragraph()
    _set_font(p.add_run(text), 13, True, (0x2C, 0x4A, 0x75), '黑体')
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(10), Pt(5)

def h3(text):
    p = doc.add_paragraph()
    _set_font(p.add_run(text), 12, True, (0x44, 0x44, 0x44), '黑体')
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(6), Pt(3)

def p(text, bold=False, size=12, indent=True):
    para = doc.add_paragraph()
    if indent:
        para.paragraph_format.first_line_indent = Pt(24)
    _set_font(para.add_run(text), size, bold)
    para.paragraph_format.space_after = Pt(4)
    para.paragraph_format.line_spacing = 1.3

def bullet(text, bold_prefix=None):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Pt(18)
    para.paragraph_format.space_after = Pt(3)
    para.paragraph_format.line_spacing = 1.3
    if bold_prefix:
        _set_font(para.add_run('· ' + bold_prefix), 12, True)
        _set_font(para.add_run(text), 12)
    else:
        _set_font(para.add_run('· ' + text), 12)

def table(headers, rows, widths=None, fontsize=10.5):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = 'Table Grid'
    for j, htxt in enumerate(headers):
        cell = t.rows[0].cells[j]
        cell.text = ''
        _set_font(cell.paragraphs[0].add_run(htxt), fontsize, True, name='黑体')
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.rows[i + 1].cells[j]
            cell.text = ''
            _set_font(cell.paragraphs[0].add_run(str(val)), fontsize)
    if widths:
        for j, w in enumerate(widths):
            for row in t.rows:
                row.cells[j].width = Mm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

# ============ 封面 ============
tp = doc.add_paragraph(); tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
_set_font(tp.add_run('纯视觉道路线形语义参数估计研究（课题甲 v2）'), 20, True, (0x1F, 0x3A, 0x5F), '黑体')
sp = doc.add_paragraph(); sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
_set_font(sp.add_run('—— SCI 论文题目 · 重构大纲 · 技术路线 · 可行性 · 硕士论文大纲 · 学生实施计划 ——'), 13, True, (0x2C, 0x4A, 0x75), '黑体')
sp2 = doc.add_paragraph(); sp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
_set_font(sp2.add_run('核心创新：弧长参数化公路三维线形模型（博士论文）作为表示层/先验，注入纯视觉识别，输出设计语义参数（R/A/i/E）并接车辆动力学预判；3DGS 作对照/增强载体'), 10.5, False, (0x66, 0x66, 0x66))
mp = doc.add_paragraph(); mp.alignment = WD_ALIGN_PARAGRAPH.CENTER
_set_font(mp.add_run('2026 年 9 月制定 · 面向 0 基础硕士研究生 · 周期 2026.09–2027.09'), 10.5, False, (0x99, 0x99, 0x99))

# ============ 一、核心创新与可行性（skills 分析） ============
h1('一、核心创新定稿与可行性分析（skills 深度分析）')

h2('1.1 核心创新一句话')
p('把博士论文《公路几何设计三维线形评价模型研究》的「基于弧长积分的公路三维几何模型」（准线-直母线直纹面；平面=直线/圆曲线/回旋线按桩号弧长分段积分，纵断面=直线+二次抛物线竖曲线，横断面=超高直母线；Darboux 标架几何不变量 κg/κn/τg）作为表示层与先验，注入纯视觉识别（单目/双目+IMU/轮速尺度锚定，无 LiDAR/高精地图/RTK），把视觉几何观测精化为**设计语义参数**（半径 R、回旋线参数 A、纵坡 i、超高 E(s)、曲率/坡度/挠率序列），并以车辆动力学模型验证其对运行安全舒适性预判的合理性。', bold=True)

h2('1.2 赛道判定：语义参数估计赛道 vs 重建赛道（不构成创新性撞车）')
table(['方法', '赛道', '输出', '局限（本方案补的）'], [
    ['CurveGaussian (ICCV 2025)', '重建（通用轮廓曲线）', '贝塞尔曲线坐标', '参数无工程语义；无规范约束；无动力学；输入为边缘图'],
    ['RoGS (arXiv 2405.14342)', '重建（道路路面网格）', '高程网格', '曲率后处理；无设计先验/超高/纵坡语义；无动力学验证'],
    ['GTLR-GS (arXiv 2603.23192)', '重建（室内 LiDAR 正则）', '渲染优化几何', '曲率细分是渲染导向；非道路场景'],
    ['LaneCPP (CVPR 2024)', '感知（3D 车道线）', '连续 3D 曲线', '通用物理先验；不输出设计要素参数；无文件级先验'],
    ['**本方案**', '**线形语义参数估计**', '**R/A/i/E(s)+κ/G/τ 序列**', '**工程语义输出+JTG D20 约束+车辆动力学预判（他法所无）**'],
], widths=[40, 32, 32, 58])
p('结论：现有方法全部是"重建导向"——输出更准的几何；本方案是"语义导向"——把几何变成可评价、可预判驾驶行为的工程参数。部件级先例（3DGS/曲线重建/曲率自适应）不影响本方案的组合创造性，但开题必须写差异化声明：对 CurveGaussian/RoGS/GTLR-GS 逐条"重建导向 vs 语义导向"。', size=11)

h2('1.3 三实证清单（无实证=claim，写进消融/验证章）')
bullet('E1 精化增益：同一输入下「视觉/3DGS 重建 + 模型精化」vs「直接重建/后处理求曲率」，R/A/i 设计参数精度提升（消融对比）。')
bullet('E2 可辨识性：含噪几何反演模型参数的唯一性与初值敏感性（混合优化：要素切分+段参数）；开题前 2~4 周数值预研，三问：可辨识性/初值鲁棒性/相对样条多项式基线增益。')
bullet('E3 语义/物理合理性：估计参数 JTG D20 合规率；设计-施工偏差检出能力；输入滚动-振动模型后与实测振动/受力吻合度（车辆动力学预判合理性）。')

h2('1.4 风险评级与对策')
table(['风险', '等级', '对策'], [
    ['大类先例（参数化+物理先验注入视觉）', '中', 'LaneCPP 差异化声明；锁"设计语义输出"层'],
    ['3DGS 部件撞车（CurveGaussian/RoGS/GTLR-GS）', '中', '3DGS 降级为对照/增强章+展望，不作核心必须部件；若进主干须逐条差异化'],
    ['要素切分+参数估计初值/可辨识性', '中高', '开题前数值预研（E2），结论决定档位'],
    ['0 基础学生 DL 训练门槛（RoadBEV 微调）', '中', '基础期 PyTorch 集训前置；训练范围限微调'],
    ['单目尺度模糊 / 前瞻距离', '中', '尺度锚定（车高/轮速/双目）；流式多帧拼接 5m→50m'],
], widths=[62, 14, 86])
p('总体评级：B+~A-（三实证齐备可冲一档期刊）。', size=12, bold=True)

# ============ 二、SCI 论文 ============
h1('二、高水平 SCI 论文（题目 + 大纲 + 投稿定位）')

h2('2.1 推荐题目（主推）')
tp2 = doc.add_paragraph(); tp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
_set_font(tp2.add_run('Vision-based Road Curvature and Grade Estimation via Arc-Length-Parameterized Highway Alignment Model: Engineering-Semantic Representation and Vehicle-Dynamics Validation'), 11.5, True, (0x1F, 0x3A, 0x5F), 'Times New Roman')
tp3 = doc.add_paragraph(); tp3.alignment = WD_ALIGN_PARAGRAPH.CENTER
_set_font(tp3.add_run('中文：基于弧长参数化公路三维线形模型的纯视觉道路曲率与坡度估计：工程语义表示与车辆动力学验证'), 12, True, (0x2C, 0x4A, 0x75), '黑体')

h2('2.2 论文大纲（SCI 六章）')
h3('Chapter 1 Introduction')
bullet('Background: 去图化/众包自动驾驶；道路几何对安全舒适影响；悬架预瞄与线形复核需求；底座 C 端。')
bullet('Related work 四线+差距表：视觉坡度/曲率（IET-ITS 2019）；3D 车道检测与先验融合（LaneCPP/Depth3DLane/PriorLane/P-MapNet）；道路线形提取重建（CACIE 2020/Springer 2023）；路面高程/3DGS 重建（RoadBEV T-ITS 2024/RoGS/CurveGaussian）。')
bullet('Contribution（三创新）：①弧长参数化线形模型表示层；②模型-数据融合精化；③车辆动力学验证。')
h3('Chapter 2 Arc-Length-Parameterized Highway Alignment Model（创新①）')
bullet('2.1 直纹面曲面模型：准线（路线设计线）+直母线（横坡/超高方向），平面/纵断面/横断面分段描述（直线/圆曲线/回旋线/竖曲线二次抛物线/超高线性渐变）。')
bullet('2.2 Darboux 标架几何不变量：测地曲率 κg、法曲率 κn、测地挠率 τg 的弧长积分表达；平纵横耦合（圆曲线×纵坡→空间螺旋线）。')
bullet('2.3 规范约束体系：G1/G2 连续、最大离心加速度、转向角/转向角速度限制、超高上限、爬坡能力（博士论文第二章约束体系）。')
bullet('2.4 估计问题定义与可辨识性分析（E2 理论部分）。')
h3('Chapter 3 Vision Observation and Model-Data Fusion Refinement（创新②）')
bullet('3.1 视觉几何观测：RoadBEV 高程场（主链路，单目/双目）+ 车道线/深度辅助；3DGS 重建对照（RoGS 开源，对照章）。')
bullet('3.2 尺度锚定：已知车高基准面/轮速积分/双目视差定量对比。')
bullet('3.3 要素切分与参数估计：直线/圆曲线/回旋线/竖曲线/超高段的联合切分与 R/A/i/E 估计（混合优化：初值策略+全局搜索+局部精化）。')
bullet('3.4 约束拟合：数据项+平滑+G²连续+JTG D20 规范+GE 域设计先验容差带；流式多帧拼接与桩号锚定（5m→50m 前瞻）。')
h3('Chapter 4 Experiments: Semantic Parameter Accuracy（E1/E2）')
bullet('4.1 RSRD 定量：R/A/i 参数误差 vs 点云 GT；「视觉+模型精化」vs「直接重建/后处理」消融（约束逐项移除）。')
bullet('4.2 3DGS 对照：RoGS 重建→同一精化管道，载体增益对比。')
bullet('4.3 初值鲁棒性：多初值/噪声水平下的参数收敛统计（E2 实验）。')
bullet('4.4 G228 实拍：估计线形 vs GE 域设计线形；竣工-设计偏差反演。')
h3('Chapter 5 Vehicle-Dynamics Validation and Applications（创新③/E3）')
bullet('5.1 Darboux 不变量→速度规划/滚动-振动 1/4 车模型：识别误差→振动/受力误差传播分析。')
bullet('5.2 安全舒适预判合理性：与实测振动/受力对比；驾驶安全风险评价模型衔接（博士论文第五章）。')
bullet('5.3 工程应用：JTG D20 合规率评估、设计-施工偏差检出、线形评价/悬架预瞄展望。')
h3('Chapter 6 Conclusion and Future Work')
bullet('结论；展望：3DGS 端到端可微约束（产业版）、多车众包协同、底座 K 域回流。')

h2('2.3 投稿定位')
table(['档位', '目标', '条件'], [
    ['冲', 'IEEE T-ITS / TRC / CACIE（IF≈8-11）', '三实证齐备 + RSRD 定量赢基线 + 消融完整'],
    ['稳', 'IEEE T-IV / Automation in Construction / ESWA', '应用叙事强化（竣工复核+合规率）'],
    ['保底', '中国公路学报 / 交通运输工程学报 / IV·ITSC 会议', '预研判据不理想时启用'],
], widths=[18, 80, 64])

# ============ 三、技术路线图 ============
h1('三、技术路线图（v2）')
p('（深色底图；M2/M3 为核心创新，V3 为特色验证；图中指标为目标值）', size=10.5)
try:
    doc.add_picture('/data/cy/shujuku/output/技术路线图-课题甲-纯视觉线形估计-v2.png', width=Mm(158))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
except Exception as e:
    p('图片插入失败: %s' % e, size=10)

# ============ 四、硕士论文大纲 ============
h1('四、硕士学位论文大纲（正式版）')
p('题目：《基于弧长参数化公路三维线形模型的纯视觉道路曲率与坡度估计研究》', size=11)
h3('第一章 绪论')
bullet('1.1 研究背景与意义；1.2 国内外研究现状（四线+差距）；1.3 研究内容与技术路线；1.4 论文组织结构。')
h3('第二章 弧长参数化公路三维线形模型与语义表示')
bullet('2.1 公路三维曲面（直纹面）模型；2.2 平面/纵断面/横断面分段参数化；2.3 Darboux 标架几何不变量；2.4 规范约束体系；2.5 估计问题定义与可辨识性。')
h3('第三章 纯视觉路面几何观测与尺度锚定')
bullet('3.1 RoadBEV 原理与复现；3.2 RSRD 训练与评测；3.3 尺度锚定方案对比；3.4 多帧拼接与桩号锚定。')
h3('第四章 模型-数据融合的线形语义参数估计')
bullet('4.1 要素切分与参数估计框架；4.2 带约束的参数化拟合；4.3 初值策略与全局-局部优化；4.4 流式在线估计；4.5 消融实验。')
h3('第五章 实验验证与应用分析')
bullet('5.1 RSRD 定量验证；5.2 3DGS 对照实验；5.3 G228 实拍与竣工复核；5.4 车辆动力学耦合验证；5.5 应用讨论。')
h3('第六章 结论与展望')
bullet('6.1 主要结论；6.2 创新点；6.3 展望。')
p('（含参考文献与附录：公式推导、优化器配置、数据处理脚本说明）', size=11)

# ============ 五、学生实施计划 ============
h1('五、学生实施计划（0 基础，2026.09–2027.09）')

h2('5.1 总体时间线')
table(['阶段', '时间', '主线任务', '里程碑/产出'], [
    ['基础期', '2026.9–10', 'Python/PyTorch 集训 + 博士论文模型学习 + RoadBEV 复现', 'RoadBEV demo 跑通；模型理解文档'],
    ['理论期 T1', '2026.11–12', '第 2 章推导 + ★E2 数值预研（三问）+ GE 域换算工具', '预研报告（决定档位）；推导文档 v1'],
    ['方法期 T2', '2027.1–4', '第 4 章要素切分+约束拟合+流式估计+消融', '2027.3 全链路演示；消融表完整'],
    ['验证期 T3', '2027.5–7', 'RSRD 定量 + 3DGS 对照 + G228 实拍 + 动力学耦合', '2027.5 中期；对比数据完整；初稿 60%'],
    ['论文期', '2027.7–9', '硕论定稿 + SCI 提炼投稿 + 预答辩', '2027.9 送审；SCI 初稿投出'],
], widths=[20, 20, 62, 60])

h2('5.2 分阶段任务详解（现在做什么 / 学什么 / 下一步）')

h3('阶段 0 · 基础期（2026.9–10，6 周）')
p('现在做什么：', bold=True)
bullet('W1-2：精读博士论文第 2 章（直纹面模型、分段要素、Darboux 标架、约束体系）与曲面滚动-振动论文；通读 RoadBEV 论文与 RSRD 说明；建术语表。')
bullet('W3-4：PyTorch 集训（张量/DataLoader/训练循环/CNN）；Python 数值计算；Linux/Git；微分几何入门（曲线曲率/挠率、Frenet/Darboux 标架）。')
bullet('W5-6：复现 RoadBEV 推理（RSRD 测试集跑通官方 checkpoint，可视化高程/误差图）；用 Mathematica/scipy 复算博士论文 2.5 节几何描述公式。')
p('学什么：', bold=True)
bullet('PyTorch 深度学习基础；微分几何（够用即可）；相机模型与 BEV 投影；公路线形设计要素（JTG D20 平纵横）。')
p('验收：RoadBEV demo 跑通；公式复算通过。', bold=True)

h3('阶段 1 · 理论期 T1（2026.11–12）★ 开题前判据')
p('现在做什么：', bold=True)
bullet('第 2 章核心推导：直纹面参数化→Darboux 不变量→R/A/i/E 映射公式（符号+数值验证）。')
bullet('★ E2 数值预研（论文立论判据，2~4 周）：在 RSRD 高程/仿真数据上做「要素切分+参数估计」实验，三问：①含噪几何下要素分段与 R/A/i 能否恢复（可辨识性）；②多初值是否收敛同一解（初值鲁棒性）；③vs 样条/多项式基线精度增益。结论三档：可辨识+增益→冲一档；要素可切分但参数不稳→竣工反演叙事二档；仅符号/量级→悬架预瞄应用。')
bullet('GE 域换算工具 v1：G228 设计文件→κ(s)/G(s)/E(s) 函数库（与课题乙共用）。')
p('学什么：', bold=True)
bullet('数值优化（带约束 SQP/全局搜索）；混合整数优化概念（切分点）；scipy.optimize/PyTorch 自动微分。')
p('验收：预研报告（三问结论）；推导文档 v1；GE 工具 v1 入库。', bold=True)

h3('阶段 2 · 方法期 T2（2027.1–4）核心实现')
p('现在做什么：', bold=True)
bullet('1 月：RoadBEV 微调（RSRD 全量）；多帧拼接模块 v1。')
bullet('2 月：要素切分+参数估计 v1（初值策略+全局-局部优化），RSRD 跑通并记录收敛行为。')
bullet('3 月：带约束拟合完整版（数据+平滑+G²+JTG D20+设计先验）；流式桩号锚定+状态传播；全链路演示（2027.3）。')
bullet('4 月：消融实验（约束逐项移除：精度/平滑度/合规性对比）；初值鲁棒性统计（E2 实验）。')
p('学什么：', bold=True)
bullet('优化理论与可微编程；状态估计基础（可选）；实验方法论（消融/指标/误差传播）。')
p('验收：全链路可演示；消融表完整；每周 Git 提交。', bold=True)

h3('阶段 3 · 验证期 T3（2027.5–7）')
p('现在做什么：', bold=True)
bullet('5 月：RSRD 定量收尾（R/A/i 误差 vs 点云 GT；vs 直接重建/后处理基线）；中期检查材料。')
bullet('6 月：3DGS 对照（RoGS 开源复现→同一精化管道，载体增益对比）；G228 实拍采集与验证（估计 vs 设计；竣工-偏差）。')
bullet('7 月：车辆动力学耦合（Darboux 不变量→速度规划/滚动-振动模型，识别误差→振动/受力误差传播，与实测对比）；论文初稿 60%。')
p('验收：对比实验完整；E3 结论明确；初稿 60%。', bold=True)

h3('阶段 4 · 论文期（2027.7–9）')
p('硕论全稿与修改；SCI 提炼与投稿（见 2.3）；预答辩（2027.9）；结题材料（实现说明/测试报告）。')

h2('5.3 考核与节奏')
table(['机制', '要求'], [
    ['契约先行', '第 2 章公式、数据集接口、GE 域函数库为唯一事实源；变更走工单'],
    ['两周里程碑', '每两周可演示/可量化进展；每周 30 分钟例会'],
    ['Git 日志', 'commit 即开发日志；预研/推导/实验脚本全部入库'],
    ['写作训练', '每阶段 1 份技术文档（预研报告/推导/消融报告）=论文素材'],
], widths=[30, 132])

h2('5.4 风险与对策')
table(['风险', '等级', '对策'], [
    ['E2 预研判据不通过（要素切分/参数估计不稳）', '高', '开题前暴露：降级为竣工线形反演应用叙事（二档）或悬架预瞄应用；不硬扛'],
    ['DL 入门慢（RoadBEV 训练）', '中', '基础期集训前置；训练限微调；GPU 错峰'],
    ['单目尺度模糊', '中', '尺度锚定开题前定量对比；必要时双目起步'],
    ['前瞻距离不足（ROI 5m）', '中', '第 4 章流式拼接；超时则降级多帧+滤波'],
    ['3DGS 对照章工程量', '中', 'RoGS 开源直接复现，限对照不优化；不进主干'],
    ['发表进度', '中', '预研通过后 2027.3 投会议/短文；主论文 2027.9 前投出'],
], widths=[62, 14, 86])

h2('5.5 发表计划')
table(['档位', '目标', '投稿时间', '条件'], [
    ['冲', 'IEEE T-ITS / TRC / CACIE（IF≈8-11）', '2027.6–9 主论文', 'E2 预研通过 + 消融完整 + RSRD 定量赢基线'],
    ['稳', 'IEEE T-IV / Automation in Construction / ESWA', '2027.6–9 备选', '应用叙事强化（竣工复核+合规率）'],
    ['保底', '中国公路学报 / 交通运输工程学报；IV/ITSC 会议', '2027.3 短文', '预研判据不理想时启用'],
], widths=[18, 68, 30, 46])

# ============ 六、协同 ============
h1('六、与底座及课题乙的协同')
bullet('GE 域：本课题是"设计文件→κ/G/E 函数库"第一消费方，工具与课题乙共用，底座数据组牵头。')
bullet('K 域回流：估计结果（众包线形修正）设计为底座 K 域数据源，回填后更新 GE 域先验容差带。')
bullet('课题乙：共享 GE 工具与 G228 实拍（乙用几何约束做换道轨迹，甲用几何约束做估计——各消费 H 域一半）。')
bullet('3DGS/数字孪生/悬架预瞄：论文展望章；zidongjiashi 文档 3DGS 方案作为后续专利/产业版。')

end = doc.add_paragraph(); end.alignment = WD_ALIGN_PARAGRAPH.CENTER
_set_font(end.add_run('—— 任务一经分配即冻结，照图施工，设计变更一律由导师把关 ——'), 11, False, (0x88, 0x88, 0x88))

out = '/data/cy/shujuku/output/课题甲v2-SCI论文大纲技术路线可行性硕论大纲与学生实施计划.docx'
doc.save(out)
print('OK ->', out)
